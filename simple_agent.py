import os
from typing import TypedDict
from dotenv import load_dotenv
# import anthropic
from langgraph.graph import StateGraph, END
from langgraph.checkpoint.memory import MemorySaver
from pypdf import PdfReader
from openai import OpenAI
from docx import Document as DocxDocument
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from langsmith import traceable
from langsmith.wrappers import wrap_openai

load_dotenv()

os.environ["LANGCHAIN_TRACING_V2"] = "true"
os.environ["LANGCHAIN_API_KEY"] = os.getenv("LANGSMITH_API_KEY", "")
os.environ["LANGCHAIN_PROJECT"] = os.getenv("LANGSMITH_PROJECT", "default")

INPUT_FOLDER = os.path.join(os.path.dirname(__file__), "Input")

class Document(TypedDict):
    type_of_document: str
    document_content: str

class State(TypedDict):
    list_of_documents: list[Document]
    recommendation: str




def _read_file(filepath: str) -> str:
    if filepath.lower().endswith(".pdf"):
        reader = PdfReader(filepath)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    with open(filepath, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def node_1(state: State) -> State:
    print("Node 1 executed")
    documents = []
    for filename in sorted(os.listdir(INPUT_FOLDER)):
        filepath = os.path.join(INPUT_FOLDER, filename)
        if os.path.isfile(filepath):
            content = _read_file(filepath)
            document: Document = {
                "type_of_document": filename,
                "document_content": content,
            }
            documents.append(document)
            print(f"  Read: {filename}")
    state["list_of_documents"] = documents
    return state


def node_2(state: State) -> State:
    print("Node 2 executed")
    # client = anthropic.Anthropic(api_key=os.getenv("ANTHROPIC_API_KEY"))
    # from openai import OpenAI
    # client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
    documents_section = "\n\n".join(
        f"DOCUMENT TYPE: {doc['type_of_document']}\n\nDOCUMENT CONTENT:\n{doc['document_content']}"
        for doc in state["list_of_documents"]
    )

    prompt = f"""
       You are a tax advisor analyzing a financial document to identify tax-saving opportunities.

    {documents_section}

    TASK:
    Based on the document above, identify the 5 most impactful tax-saving tips relevant to this specific document. Prioritize tips that:
    - Are directly supported by details found in the document (income sources, deductions, expenses, investments, filing status, etc.)
    - Would have meaningful financial impact, not marginal ones
    - Are actionable (something the person could realistically do or claim)

    OUTPUT FORMAT:
    Return exactly 5 bullet points, each structured as:
    - **[Short tip title]**: [1-2 sentence explanation of the tip and why it applies based on what's in the document]

    RULES:
    - Base tips only on what's evident or reasonably inferable from the document — don't invent numbers or facts not present.
    - If the document lacks enough detail for a full 5 relevant, specific tips, state that clearly and give general tips only to fill the gap.
    - Avoid generic advice that isn't tied to the document (e.g., don't say "contribute to a retirement account" unless the document shows relevant income/employment info that makes this apply).
    - Do not provide legal or definitive tax advice — frame tips as suggestions to explore/verify with a tax professional or current tax code
    - Keep language plain and jargon-free where possible.
        """

    # response = client.messages.create(
    #     model="claude-haiku-4-5-20251001",
    #     max_tokens=1024,
    #     messages=[{"role": "user", "content": prompt}],
    # )
    # print("Claude response:", response.content[0].text)
    # return state
    client = OpenAI()
    response = client.responses.create(
        model="gpt-4.1-mini",
        input=prompt,
    )

    # print("OpenAI response:", response.output_text)
    state["recommendation"] = response.output_text
    return state

def node_3(state: State) -> State:
    print("Node 3 executed")
    doc = DocxDocument()

    title = doc.add_heading("Tax Saving Recommendations", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    recommendation_text = state.get("recommendation", "")
    lines = [line.strip() for line in recommendation_text.splitlines() if line.strip()]

    for line in lines:
        if line.startswith("- "):
            line = line[2:]

        para = doc.add_paragraph(style="List Bullet")

        if "**" in line:
            parts = line.split("**")
            for i, part in enumerate(parts):
                if not part:
                    continue
                run = para.add_run(part)
                run.bold = (i % 2 == 1)
                run.font.size = Pt(11)
        else:
            run = para.add_run(line)
            run.font.size = Pt(11)

    output_dir = os.path.join(os.path.dirname(__file__), "output")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "Recommendation1.docx")
    doc.save(output_path)
    print(f"  Saved: {output_path}")
    return state

# building the graph
builder = StateGraph(State)

builder.add_node("node_1", node_1)
builder.add_node("node_2", node_2)
builder.add_node("node_3", node_3)

builder.set_entry_point("node_1")
builder.add_edge("node_1", "node_2")
builder.add_edge("node_2", "node_3")
builder.add_edge("node_3", END)

checkpointer = MemorySaver()
graph = builder.compile(checkpointer=checkpointer)

if __name__ == "__main__":
    config = {"configurable": {"thread_id": "thread-1"}}
    result = graph.invoke({}, config=config)
    print("Done:", result)
